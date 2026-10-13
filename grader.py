import csv, glob, os
from openpyxl import load_workbook

TOL = 0.01
key = load_workbook("key.xlsx", data_only=True)
with open("rubric.csv", encoding="utf-8") as f:
    rubric = list(csv.DictReader(f))
maks = sum(float(r["points"]) for r in rubric)

def same(a, b):
    try:
        return abs(float(a) - float(b)) <= TOL
    except (TypeError, ValueError):
        return str(a).strip().lower() == str(b).strip().lower()

rows = [["fayl", "ball", "maksimum"]]
for path in sorted(glob.glob("submissions/*.xlsx")):
    sub = load_workbook(path, data_only=True)
    ball = 0
    print(f"\n== {os.path.basename(path)} ==")
    for r in rubric:
        sh, cell, pts = r["sheet"], r["cell"], float(r["points"])
        try:
            ok = same(sub[sh][cell].value, key[sh][cell].value)
        except KeyError:
            ok = False
        if ok:
            ball += pts
        print(f"{sh}!{cell}: {'TOGRI' if ok else 'XATO'} ({pts if ok else 0}/{pts})")
    print(f"JAMI: {ball}/{maks}")
    rows.append([os.path.basename(path), ball, maks])

os.makedirs("results", exist_ok=True)
with open("results/natijalar.csv", "w", newline="", encoding="utf-8") as f:
    csv.writer(f).writerows(rows)
